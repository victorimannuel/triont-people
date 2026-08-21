import io
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from extensions import db
from models.user import User, Department
from models.company import Company
from models.leave import LeaveRequest, LeaveBalance, LeaveType

def generate_leave_report_excel(company_id, year=None, department_id=None):
    """
    Generates a multi-sheet Excel (.xlsx) report containing:
    1. Leave Balance Summary (Rekap Saldo Cuti) per employee with auto-calculated formulas.
    2. Leave Request Details (Rincian Pengajuan Cuti) with complete history.
    """
    this_year = year or date.today().year
    company = db.session.get(Company, company_id)
    company_name = company.name if company else "Company"

    wb = openpyxl.Workbook()
    # Default sheet
    ws_balance = wb.active
    ws_balance.title = f"Rekap Saldo {this_year}"

    # Styling definitions
    header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    sub_header_fill = PatternFill(start_color="14B8A6", end_color="14B8A6", fill_type="solid")
    accent_fill = PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid")
    zebra_fill = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")
    total_fill = PatternFill(start_color="E6FFFA", end_color="E6FFFA", fill_type="solid")

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=15, bold=True, color="0F766E")
    meta_font = Font(name="Calibri", size=10, italic=True, color="555555")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    double_bottom_border = Border(
        left=thin_border_side, right=thin_border_side, top=thin_border_side,
        bottom=Side(border_style="double", color="0F766E")
    )

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    # ── SHEET 1: REKAP SALDO CUTI ──
    # Title block
    ws_balance.merge_cells("A1:G1")
    ws_balance["A1"] = f"LAPORAN REKAPITULASI SALDO CUTI KARYAWAN - {this_year}"
    ws_balance["A1"].font = title_font
    ws_balance["A1"].alignment = left_align

    ws_balance["A2"] = f"Perusahaan: {company_name} | Tanggal Export: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws_balance["A2"].font = meta_font

    # Active leave types
    leave_types = LeaveType.query.filter_by(company_id=company_id, is_active=True).order_by(LeaveType.id.asc()).all()

    # Query employees
    emp_query = User.query.filter_by(company_id=company_id, is_active=True, is_deleted=False)
    if department_id:
        emp_query = emp_query.filter_by(department_id=department_id)
    employees = emp_query.order_by(User.name.asc()).all()
    emp_ids = [e.id for e in employees]

    # Pre-fetch balances
    balances = LeaveBalance.query.filter(
        LeaveBalance.company_id == company_id,
        LeaveBalance.year == this_year,
        LeaveBalance.employee_id.in_(emp_ids)
    ).all()
    bal_map = {(b.employee_id, b.leave_type_id): b for b in balances}

    # Build Header Rows (Row 4 & 5)
    base_headers = ["No", "Nama Karyawan", "Email", "Departemen", "Atasan", "Role"]
    
    # Row 4 merges
    for col_idx, h in enumerate(base_headers, 1):
        ws_balance.merge_cells(start_row=4, start_column=col_idx, end_row=5, end_column=col_idx)
        cell = ws_balance.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = cell_border
        ws_balance.cell(row=5, column=col_idx).border = cell_border

    curr_col = len(base_headers) + 1
    lt_col_ranges = []

    for lt in leave_types:
        start_c = curr_col
        end_c = curr_col + 3
        ws_balance.merge_cells(start_row=4, start_column=start_c, end_row=4, end_column=end_c)
        lt_cell = ws_balance.cell(row=4, column=start_c, value=lt.name)
        lt_cell.fill = sub_header_fill
        lt_cell.font = header_font
        lt_cell.alignment = center_align
        lt_cell.border = cell_border

        sub_cols = ["Hak Cuti", "Terpakai", "Pending", "Sisa Saldo"]
        for i, sub_title in enumerate(sub_cols):
            c_idx = start_c + i
            sub_cell = ws_balance.cell(row=5, column=c_idx, value=sub_title)
            sub_cell.fill = header_fill
            sub_cell.font = header_font
            sub_cell.alignment = center_align
            sub_cell.border = cell_border
            ws_balance.cell(row=4, column=c_idx).border = cell_border

        lt_col_ranges.append((lt, start_c, end_c))
        curr_col += 4

    # Grand Total Columns
    tot_start = curr_col
    tot_end = curr_col + 3
    ws_balance.merge_cells(start_row=4, start_column=tot_start, end_row=4, end_column=tot_end)
    tot_cell = ws_balance.cell(row=4, column=tot_start, value="TOTAL KESELURUHAN")
    tot_cell.fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    tot_cell.font = header_font
    tot_cell.alignment = center_align

    tot_subs = ["Total Hak", "Total Terpakai", "Total Pending", "Total Sisa"]
    for i, sub_title in enumerate(tot_subs):
        c_idx = tot_start + i
        sub_cell = ws_balance.cell(row=5, column=c_idx, value=sub_title)
        sub_cell.fill = PatternFill(start_color="115E59", end_color="115E59", fill_type="solid")
        sub_cell.font = header_font
        sub_cell.alignment = center_align
        sub_cell.border = cell_border
        ws_balance.cell(row=4, column=c_idx).border = cell_border

    # Data Rows (Row 6 onwards)
    data_start_row = 6
    for row_num, emp in enumerate(employees, start=data_start_row):
        is_zebra = (row_num % 2 == 1)
        row_fill = zebra_fill if is_zebra else PatternFill(fill_type=None)

        dept_name = emp.department.name if emp.department else "-"
        mgr_name = emp.manager.name if emp.manager else "-"

        # Base fields
        row_values = [
            row_num - data_start_row + 1,
            emp.name,
            emp.email,
            dept_name,
            mgr_name,
            emp.role.upper()
        ]
        for c_idx, val in enumerate(row_values, 1):
            cell = ws_balance.cell(row=row_num, column=c_idx, value=val)
            cell.font = regular_font
            cell.alignment = center_align if c_idx == 1 or c_idx == 6 else left_align
            if is_zebra:
                cell.fill = row_fill
            cell.border = cell_border

        # Leave types values
        tot_quota_cells = []
        tot_used_cells = []
        tot_pend_cells = []
        tot_rem_cells = []

        for lt, start_c, _ in lt_col_ranges:
            bal = bal_map.get((emp.id, lt.id))
            quota = float(bal.total_days) if bal else (float(lt.days_per_year) if lt.days_per_year else 0.0)
            used = float(bal.used_days) if bal else 0.0
            pending = float(bal.pending_days) if bal else 0.0
            
            c_quota = get_column_letter(start_c)
            c_used = get_column_letter(start_c + 1)
            c_pend = get_column_letter(start_c + 2)
            c_rem = get_column_letter(start_c + 3)

            # Cell 1: Quota
            cell_q = ws_balance.cell(row=row_num, column=start_c, value=quota)
            cell_q.font = regular_font
            cell_q.alignment = right_align
            cell_q.number_format = "#,##0.0"
            cell_q.border = cell_border

            # Cell 2: Used
            cell_u = ws_balance.cell(row=row_num, column=start_c + 1, value=used)
            cell_u.font = regular_font
            cell_u.alignment = right_align
            cell_u.number_format = "#,##0.0"
            cell_u.border = cell_border

            # Cell 3: Pending
            cell_p = ws_balance.cell(row=row_num, column=start_c + 2, value=pending)
            cell_p.font = regular_font
            cell_p.alignment = right_align
            cell_p.number_format = "#,##0.0"
            cell_p.border = cell_border

            # Cell 4: Formula Sisa = Quota - Used - Pending
            cell_r = ws_balance.cell(row=row_num, column=start_c + 3, value=f"={c_quota}{row_num}-{c_used}{row_num}-{c_pend}{row_num}")
            cell_r.font = bold_font
            cell_r.alignment = right_align
            cell_r.number_format = "#,##0.0"
            cell_r.border = cell_border

            tot_quota_cells.append(f"{c_quota}{row_num}")
            tot_used_cells.append(f"{c_used}{row_num}")
            tot_pend_cells.append(f"{c_pend}{row_num}")
            tot_rem_cells.append(f"{c_rem}{row_num}")

            if is_zebra:
                cell_q.fill = row_fill
                cell_u.fill = row_fill
                cell_p.fill = row_fill
                cell_r.fill = row_fill

        # Totals for employee row (Formulas)
        t_q_cell = ws_balance.cell(row=row_num, column=tot_start, value=f"=SUM({'+'.join(tot_quota_cells)})" if tot_quota_cells else 0)
        t_u_cell = ws_balance.cell(row=row_num, column=tot_start + 1, value=f"=SUM({'+'.join(tot_used_cells)})" if tot_used_cells else 0)
        t_p_cell = ws_balance.cell(row=row_num, column=tot_start + 2, value=f"=SUM({'+'.join(tot_pend_cells)})" if tot_pend_cells else 0)
        t_r_cell = ws_balance.cell(row=row_num, column=tot_start + 3, value=f"=SUM({'+'.join(tot_rem_cells)})" if tot_rem_cells else 0)

        for i, c in enumerate([t_q_cell, t_u_cell, t_p_cell, t_r_cell]):
            c.font = bold_font
            c.alignment = right_align
            c.number_format = "#,##0.0"
            c.fill = accent_fill
            c.border = cell_border

    # Bottom Summary Row (SUM across all employees)
    if employees:
        last_emp_row = data_start_row + len(employees) - 1
        summary_row = last_emp_row + 1

        ws_balance.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=6)
        sum_label = ws_balance.cell(row=summary_row, column=1, value="TOTAL KESELURUHAN (HARI)")
        sum_label.font = bold_font
        sum_label.alignment = Alignment(horizontal="right", vertical="center")
        sum_label.fill = total_fill

        for c in range(1, 7):
            ws_balance.cell(row=summary_row, column=c).border = double_bottom_border
            ws_balance.cell(row=summary_row, column=c).fill = total_fill

        # Compute SUM formula for each numeric column
        max_col = tot_end
        for c_idx in range(7, max_col + 1):
            col_letter = get_column_letter(c_idx)
            cell = ws_balance.cell(row=summary_row, column=c_idx, value=f"=SUM({col_letter}{data_start_row}:{col_letter}{last_emp_row})")
            cell.font = bold_font
            cell.alignment = right_align
            cell.number_format = "#,##0.0"
            cell.fill = total_fill
            cell.border = double_bottom_border

    # Auto-adjust column widths for Sheet 1
    for col in ws_balance.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if cell.number_format and ('=' not in val_str):
                val_str = str(cell.value)
            if len(val_str) > max_len and cell.row > 2:
                max_len = len(val_str)
        ws_balance.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws_balance.column_dimensions["A"].width = 6
    ws_balance.column_dimensions["B"].width = 24
    ws_balance.column_dimensions["C"].width = 28
    ws_balance.column_dimensions["D"].width = 20
    ws_balance.column_dimensions["E"].width = 22

    # ── SHEET 2: RINCIAN PENGAJUAN CUTI ──
    ws_details = wb.create_sheet(title=f"Rincian Pengajuan {this_year}")

    ws_details.merge_cells("A1:G1")
    ws_details["A1"] = f"DAFTAR RINCIAN PENGAJUAN CUTI KARYAWAN - {this_year}"
    ws_details["A1"].font = title_font
    ws_details["A1"].alignment = left_align

    ws_details["A2"] = f"Perusahaan: {company_name} | Tanggal Export: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws_details["A2"].font = meta_font

    # Detail Headers
    detail_headers = [
        ("No", 6, center_align),
        ("ID Pengajuan", 14, center_align),
        ("Nama Karyawan", 24, left_align),
        ("Email", 28, left_align),
        ("Departemen", 20, left_align),
        ("Jenis Cuti", 20, left_align),
        ("Tanggal Mulai", 15, center_align),
        ("Tanggal Selesai", 15, center_align),
        ("Tipe Durasi", 14, center_align),
        ("Durasi (Hari)", 14, right_align),
        ("Status", 16, center_align),
        ("Alasan", 32, left_align),
        ("Lampiran", 20, left_align),
        ("Delegate / Backup", 22, left_align),
        ("Tgl Pengajuan", 18, center_align),
    ]

    for col_idx, (h_name, width, align) in enumerate(detail_headers, 1):
        cell = ws_details.cell(row=4, column=col_idx, value=h_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = cell_border
        ws_details.column_dimensions[get_column_letter(col_idx)].width = width

    # Query Leave Requests for this company and year
    req_query = LeaveRequest.query.filter(
        LeaveRequest.company_id == company_id,
        db.extract('year', LeaveRequest.start_date) == this_year
    )
    if department_id:
        req_query = req_query.join(User, LeaveRequest.employee_id == User.id).filter(User.department_id == department_id)
    
    requests = req_query.order_by(LeaveRequest.created_at.desc()).all()

    for row_idx, r in enumerate(requests, start=5):
        is_zebra = (row_idx % 2 == 1)
        row_fill = zebra_fill if is_zebra else PatternFill(fill_type=None)

        day_part_label = "Seharian (Full)" if r.day_part == 'full' else ("Pagi (0.5)" if r.day_part == 'morning' else "Siang (0.5)")
        status_label = "Disetujui (Approved)" if r.status == 'approved' else ("Ditolak (Rejected)" if r.status == 'rejected' else "Menunggu (Pending)")
        attachment_label = r.attachment_original_name or ("Ada Lampiran" if r.attachment_path else "-")
        delegate_name = r.delegate.name if r.delegate else "-"
        created_str = r.created_at.strftime('%d/%m/%Y %H:%M') if r.created_at else "-"

        vals = [
            row_idx - 4,
            f"REQ-{r.id:04d}",
            r.employee.name if r.employee else "-",
            r.employee.email if r.employee else "-",
            r.employee.department.name if (r.employee and r.employee.department) else "-",
            r.leave_type.name if r.leave_type else "-",
            r.start_date.strftime('%d/%m/%Y') if r.start_date else "-",
            r.end_date.strftime('%d/%m/%Y') if r.end_date else "-",
            day_part_label,
            float(r.duration_days) if r.duration_days else 0.0,
            status_label,
            r.reason or "-",
            attachment_label,
            delegate_name,
            created_str,
        ]

        for col_idx, (val, _, align) in enumerate(zip(vals, [h[1] for h in detail_headers], [h[2] for h in detail_headers]), 1):
            cell = ws_details.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.alignment = align
            cell.border = cell_border
            if is_zebra:
                cell.fill = row_fill
            if col_idx == 10:  # Duration
                cell.number_format = "#,##0.0"

    # Save to in-memory bytes buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_audit_logs_excel(company_id, logs):
    """
    Generates an Excel (.xlsx) file containing audit logs.
    """
    import json
    company = db.session.get(Company, company_id)
    company_name = company.name if company else "Company"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit Logs"

    # Styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=15, bold=True, color="0F172A")
    meta_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    regular_font = Font(name="Calibri", size=10)
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    left_wrap_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Title block
    ws.merge_cells("A1:I1")
    t_cell = ws.cell(row=1, column=1, value=f"AUDIT LOG & JEJAK AKTIVITAS - {company_name.upper()}")
    t_cell.font = title_font
    t_cell.alignment = left_align

    ws.merge_cells("A2:I2")
    m_cell = ws.cell(row=2, column=1, value=f"Diunduh pada: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Total Record: {len(logs)}")
    m_cell.font = meta_font
    m_cell.alignment = left_align

    headers = [
        ("No", 6, center_align),
        ("Waktu (UTC)", 20, center_align),
        ("User / Aktor", 25, left_align),
        ("Role", 14, center_align),
        ("Aksi / Operasi", 30, left_align),
        ("Target", 20, left_align),
        ("Rincian Operasi", 45, left_wrap_align),
        ("IP Address", 18, center_align),
        ("User Agent", 35, left_wrap_align),
    ]

    # Header row
    for col_idx, (h_title, w, align) in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align
        cell.border = cell_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = w

    ws.row_dimensions[4].height = 26

    for row_idx, log in enumerate(logs, 5):
        ws.row_dimensions[row_idx].height = 24
        is_zebra = (row_idx % 2 == 0)
        actor_name = f"{log.actor.name} ({log.actor.email})" if log.actor else "System / Anonymous"
        actor_role = log.actor.role.upper() if log.actor else "SYSTEM"
        target_str = f"{log.target_type} #{log.target_id}" if log.target_type else "-"
        
        details_str = json.dumps(log.details_dict, ensure_ascii=False) if log.details_dict else "-"

        from core.time_util import format_user_datetime
        vals = [
            row_idx - 4,
            format_user_datetime(log.created_at, '%d/%m/%Y %H:%M:%S') if log.created_at else "-",
            actor_name,
            actor_role,
            f"{log.action_label} ({log.action})",
            target_str,
            details_str,
            log.ip_address or "-",
            log.user_agent or "-"
        ]

        for col_idx, (val, _, align) in enumerate(zip(vals, [h[1] for h in headers], [h[2] for h in headers]), 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.alignment = align
            cell.border = cell_border
            if is_zebra:
                cell.fill = zebra_fill

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

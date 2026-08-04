import os
from datetime import date, datetime, timedelta
from functools import wraps

from dotenv import load_dotenv
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, abort
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash

load_dotenv()

from models import db, User, LeaveType, LeaveRequest, LeaveBalance, Company, ApprovalConfig, Department

# ──────────────────────── App Factory ────────────────────────

def create_app():
    app = Flask(__name__)
    app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'gwcuti-secret-change-in-production')
    app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get('DATABASE_URL', 'postgresql://gwcuti:***@localhost:5432/gwcuti')
    app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

    db.init_app(app)

    login_manager = LoginManager()
    login_manager.init_app(app)
    login_manager.login_view = 'auth.login'
    login_manager.login_message = 'Silakan login terlebih dahulu.'

    @login_manager.user_loader
    def load_user(user_id):
        return db.session.get(User, int(user_id))

    with app.app_context():
        from models import User, LeaveType, LeaveRequest, LeaveBalance
        db.create_all()
        seed_initial_data()

    # ── Register blueprints (inline for simplicity) ──
    register_routes(app)

    # ── Template context processors ──
    @app.context_processor
    def inject_globals():
        return {'date': date, 'datetime': datetime}

    @app.template_filter('indonesian_date')
    def indonesian_date_filter(d):
        if not d:
            return ''
        months = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni',
                  'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']
        return f"{d.day} {months[d.month - 1]} {d.year}"

    # ── Theme color context processor ──
    @app.context_processor
    def inject_theme():
        if current_user.is_authenticated:
            company = Company.query.get(current_user.company_id)
            primary = company.primary_color if company else '#0d9488'
            return dict(primary_color=primary)
        return dict(primary_color='#0d9488')

    return app

# ──────────────────────── Notification Helper ────────────────────────

def send_notification(company, event, leave_request):
    """Send notification based on company config. Logs to console if SMTP not configured."""
    if not company or not company.notifications_enabled:
        return

    if event == 'approve' and not company.notify_on_approve:
        return
    if event == 'reject' and not company.notify_on_reject:
        return
    if event == 'submit' and not company.notify_on_submit:
        return

    emp = leave_request.employee
    status_text = {'approve': 'disetujui', 'reject': 'ditolak', 'submit': 'diajukan'}
    subject = f'Cuti {status_text[event]} — {emp.name}'
    body = f"""
Halo {emp.name},

Pengajuan cuti {leave_request.leave_type.name} Anda telah {status_text[event]}.
Tanggal: {leave_request.start_date.strftime('%d %B %Y')} — {leave_request.end_date.strftime('%d %B %Y')}
Durasi: {int(leave_request.duration_days)} hari
Status: {leave_request.status.title()}

- GwCuti
"""

    if company.smtp_host and company.smtp_user:
        try:
            import smtplib
            from email.message import EmailMessage
            msg = EmailMessage()
            msg['Subject'] = subject
            msg['From'] = company.smtp_user
            msg['To'] = emp.email
            msg.set_content(body)

            with smtplib.SMTP(company.smtp_host, company.smtp_port) as server:
                server.starttls()
                server.login(company.smtp_user, company.smtp_password)
                server.send_message(msg)
            print(f'✅ Notification sent to {emp.email}')
        except Exception as e:
            print(f'❌ Failed to send email: {e}')

    else:
        print('ℹ️  No SMTP config. Notification logged.')

# ──────────────────────── Seed Data ────────────────────────

def seed_initial_data():
    if Company.query.first():
        return

    # Create admin company
    admin_company = Company(name='Admin', primary_color='#0d9488', is_active=True,
                            notifications_enabled=True, smtp_host='smtp.gmail.com',
                            smtp_port=587, smtp_user='admin@example.com',
                            smtp_password='password', notify_on_approve=True,
                            notify_on_reject=True, notify_on_submit=True)
    db.session.add(admin_company)
    db.session.flush()

    # Create default employee company
    emp_company = Company(name='Employee', primary_color='#0d9488', is_active=True,
                          notifications_enabled=False)
    db.session.add(emp_company)
    db.session.flush()

    # Create leave types per company
    leave_types = [
        LeaveType(name='Cuti Tahunan', description='Cuti tahunan reguler', days_per_year=12, color='#0d9488', is_active=True),
        LeaveType(name='Cuti Sakit', description='Cuti karena sakit', days_per_year=0, color='#ef4444', is_active=True),
        LeaveType(name='Cuti Melahirkan', description='Cuti melahirkan', days_per_year=90, color='#ec4899', is_active=True),
    ]
    for company in [admin_company, emp_company]:
        for lt in leave_types:
            leave_type = LeaveType(name=lt.name, description=lt.description, days_per_year=lt.days_per_year,
                                   color=lt.color, is_active=True, company_id=company.id)
            db.session.add(leave_type)
    db.session.flush()

    # Create users
    users_data = [
        {'name': 'Admin User', 'email': 'admin@gwcuti.app', 'role': 'admin', 'company_id': admin_company.id},
        {'name': 'Manager User', 'email': 'manager@gwcuti.app', 'role': 'manager', 'company_id': admin_company.id},
        {'name': 'HR User', 'email': 'hr@gwcuti.app', 'role': 'hr', 'company_id': admin_company.id},
        {'name': 'Employee User', 'email': 'employee@gwcuti.app', 'role': 'employee', 'company_id': emp_company.id},
        {'name': 'Bob Sunshine', 'email': 'bob@gwcuti.app', 'role': 'manager', 'company_id': emp_company.id},
    ]
    for user_data in users_data:
        user = User(name=user_data['name'], email=user_data['email'], role=user_data['role'],
                     company_id=user_data['company_id'], department_id=None)
        user.set_password(user_data['email'].split('@')[1] + '123')
        db.session.add(user)
    db.session.flush()

    # Create manager/subordinate relation
    employee_user = User.query.filter_by(email='employee@gwcuti.app').first()
    bob_user = User.query.filter_by(email='bob@gwcuti.app').first()
    if employee_user and bob_user:
        employee_user.manager_id = bob_user.id

    # Create departments
    admin_dept = Department(name='IT', company_id=admin_company.id)
    db.session.add(admin_dept)
    dept = Department(name='Sales', company_id=emp_company.id)
    db.session.add(dept)

    db.session.commit()

# ──────────────────────── Role Decorator ────────────────────────

def role_required(*roles):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if not current_user.is_authenticated or current_user.role not in roles:
                flash('Akses ditolak.', 'danger')
                return redirect(url_for('main.dashboard'))
            return f(*args, **kwargs)
        return decorated_function
    return decorator

# ──────────────────────── Routes ────────────────────────

def register_routes(app):
    from flask import Blueprint

    auth_bp = Blueprint('auth', __name__, url_prefix='/auth')
    main_bp = Blueprint('main', __name__)

    # ── AUTH ──

    @auth_bp.route('/login', methods=['GET', 'POST'])
    def login():
        if current_user.is_authenticated:
            return redirect(url_for('main.dashboard'))
        if request.method == 'POST':
            email = request.form.get('email', '').strip()
            password = request.form.get('password', '')
            user = User.query.filter_by(email=email).first()
            if user and user.check_password(password) and user.is_active:
                login_user(user)
                next_page = request.args.get('next')
                flash(f'Selamat datang, {user.name}!', 'success')
                return redirect(next_page or url_for('main.dashboard'))
            flash('Email atau password salah.', 'danger')
        return render_template('login.html')

    @auth_bp.route('/logout')
    @login_required
    def logout():
        logout_user()
        flash('Anda telah logout.', 'info')
        return redirect(url_for('auth.login'))

    app.register_blueprint(auth_bp)

    # ── DASHBOARD ──

    @main_bp.route('/')
    @login_required
    def dashboard():
        this_year = date.today().year
        my_company = current_user.company_id

        # My leave balances
        balances = LeaveBalance.query.filter_by(employee_id=current_user.id, year=this_year).all()
        balance_data = []
        for bal in balances:
            lt = LeaveType.query.get(bal.leave_type_id)
            if lt:
                remaining = bal.total_days - bal.used_days - bal.pending_days
                balance_data.append({
                    'id': lt.id,
                    'name': lt.name,
                    'color': lt.color,
                    'icon': lt.icon,
                    'total': bal.total_days,
                    'used': bal.used_days,
                    'pending': bal.pending_days,
                    'remaining': max(0, remaining),
                })

        # My recent requests
        my_requests = LeaveRequest.query.filter_by(employee_id=current_user.id)\
            .order_by(LeaveRequest.created_at.desc()).limit(5).all()

        # Pending approvals (for managers/admins)
        pending_approvals = []
        if current_user.role in ('manager', 'hr', 'admin'):
            if current_user.role == 'manager':
                pending_approvals = LeaveRequest.query.filter(
                    LeaveRequest.status == 'pending',
                    LeaveRequest.company_id == my_company,
                    LeaveRequest.current_approval_level == 1,
                    LeaveRequest.employee_id.in_(
                        db.session.query(User.id).filter(User.manager_id == current_user.id)
                    )
                ).order_by(LeaveRequest.created_at.desc()).all()
            else:
                pending_approvals = LeaveRequest.query.filter_by(status='pending', company_id=my_company)\
                    .order_by(LeaveRequest.created_at.desc()).all()

        # Get all companies for admin chart filter
        companies = []
        if current_user.role == 'admin':
            companies = [{'id': c.id, 'name': c.name} for c in Company.query.filter_by(is_active=True).all()]

        return render_template('dashboard.html',
            balances=balance_data,
            my_requests=my_requests,
            pending_approvals=pending_approvals,
            this_year=this_year,
            companies=companies,
            current_company=my_company)

    # ── APPLY LEAVE ──

    @main_bp.route('/apply', methods=['GET', 'POST'])
    @login_required
    @role_required('employee', 'manager', 'hr')
    def apply_leave():
        this_year = date.today().year
        my_company = current_user.company_id
        leave_types = LeaveType.query.filter_by(is_active=True, company_id=my_company).all()
        balances = {b.leave_type_id: b for b in LeaveBalance.query.filter_by(employee_id=current_user.id, year=this_year).all()}
        max_date = date(2030, 12, 31)
        employees = User.query.filter(User.id != current_user.id, User.is_active == True, User.company_id == my_company).order_by(User.name).all()

        if request.method == 'POST':
            leave_type_id = request.form.get('leave_type', type=int)
            start_date_str = request.form.get('start_date', '')
            end_date_str = request.form.get('end_date', '')
            reason = request.form.get('reason', '')
            delegate_id = request.form.get('delegate_id', type=int)

            if not all([leave_type_id, start_date_str, end_date_str]):
                flash('Harap isi semua field yang wajib.', 'danger')
                return render_template('apply.html', leave_types=leave_types, balances=balances, max_date=max_date, employees=employees)

            try:
                start_date = date.fromisoformat(start_date_str)
                end_date = date.fromisoformat(end_date_str)
            except:
                flash('Format tanggal tidak valid.', 'danger')
                return render_template('apply.html', leave_types=leave_types, balances=balances, max_date=max_date, employees=employees)

            if end_date < start_date:
                flash('Tanggal selesai harus setelah atau sama dengan tanggal mulai.', 'danger')
                return render_template('apply.html', leave_types=leave_types, balances=balances, max_date=max_date, employees=employees)

            if start_date < date.today():
                flash('Tidak bisa mengajukan cuti di masa lalu.', 'danger')
                return render_template('apply.html', leave_types=leave_types, balances=balances, max_date=max_date, employees=employees)

            # Check overlapping leave
            overlap = LeaveRequest.query.filter(
                LeaveRequest.employee_id == current_user.id,
                LeaveRequest.status.in_(['pending', 'approved']),
                LeaveRequest.start_date <= end_date,
                LeaveRequest.end_date >= start_date
            ).first()
            if overlap:
                flash('Anda sudah memiliki pengajuan cuti di tanggal tersebut.', 'danger')
                return render_template('apply.html', leave_types=leave_types, balances=balances, max_date=max_date, employees=employees)

            # Check balance
            lt = LeaveType.query.get(leave_type_id)
            bal = balances.get(leave_type_id)
            duration = (end_date - start_date).days + 1
            if lt and bal and lt.days_per_year > 0:
                remaining = bal.total_days - bal.used_days - bal.pending_days
                if duration > remaining:
                    flash(f'Cuti {lt.name} Anda tidak mencukupi. Sisa: {int(remaining)} hari.', 'danger')
                    return render_template('apply.html', leave_types=leave_types, balances=balances, max_date=max_date, employees=employees)

            # Submit
            lt_obj = LeaveType.query.get(leave_type_id)
            approval_configs = ApprovalConfig.query.filter(
                ApprovalConfig.company_id == my_company,
                (ApprovalConfig.leave_type_id == leave_type_id) | (ApprovalConfig.leave_type_id.is_(None))
            ).order_by(ApprovalConfig.level).all()
            max_level = max((c.level for c in approval_configs if c.leave_type_id == leave_type_id),
                          default=max((c.level for c in approval_configs if c.leave_type_id is None), default=1))

            req = LeaveRequest(
                employee_id=current_user.id,
                leave_type_id=leave_type_id,
                company_id=my_company,
                start_date=start_date,
                end_date=end_date,
                duration_days=duration,
                reason=reason,
                delegate_id=delegate_id if delegate_id else None,
                status='pending',
                current_approval_level=1,
                max_approval_level=max_level
            )
            db.session.add(req)
            if bal:
                bal.pending_days += duration
            db.session.commit()

            flash('Pengajuan cuti berhasil dikirim! Menunggu persetujuan.', 'success')
            return redirect(url_for('main.dashboard'))

        return render_template('apply.html', leave_types=leave_types, balances=balances, max_date=max_date, employees=employees)

    # ── APPROVALS ──

    @main_bp.route('/approvals')
    @login_required
    @role_required('manager', 'hr', 'admin')
    def approvals():
        my_company = current_user.company_id
        if current_user.role == 'manager':
            subordinates = db.session.query(User.id).filter(User.manager_id == current_user.id).subquery()
            requests = LeaveRequest.query.filter(
                LeaveRequest.employee_id.in_(subordinates),
                LeaveRequest.company_id == my_company,
                LeaveRequest.current_approval_level == 1,
                LeaveRequest.status == 'pending'
            ).order_by(LeaveRequest.created_at.desc()).all()
        else:
            requests = LeaveRequest.query.filter_by(status='pending', company_id=my_company).order_by(LeaveRequest.created_at.desc()).all()

        all_history = LeaveRequest.query.filter(
            LeaveRequest.status != 'pending',
            LeaveRequest.company_id == my_company
        ).order_by(LeaveRequest.updated_at.desc()).limit(20).all()

        return render_template('approvals.html', pending_requests=requests, history=all_history)

    @main_bp.route('/approve/<int:request_id>', methods=['POST'])
    @login_required
    @role_required('manager', 'hr', 'admin')
    def approve_leave(request_id):
        req = LeaveRequest.query.get_or_404(request_id)

        config = ApprovalConfig.query.filter(
            ApprovalConfig.company_id == req.company_id,
            (ApprovalConfig.leave_type_id == req.leave_type_id) | (ApprovalConfig.leave_type_id.is_(None)),
            ApprovalConfig.level == req.current_approval_level
        ).order_by(ApprovalConfig.leave_type_id.desc()).first()

        if current_user.role == 'manager':
            if not config or config.approver_role != 'manager':
                abort(403)
            if req.employee.manager_id != current_user.id:
                abort(403)
        elif current_user.role == 'hr':
            if not config or config.approver_role not in ('hr', 'manager'):
                abort(403)

        action = request.form.get('action')
        notes = request.form.get('notes', '')

        if action == 'approve':
            req.approved_by = current_user.id
            req.approved_at = datetime.utcnow()
            req.notes = notes

            if req.current_approval_level >= req.max_approval_level:
                req.status = 'approved'
                this_year = req.start_date.year
                bal = LeaveBalance.query.filter_by(employee_id=req.employee_id, leave_type_id=req.leave_type_id, year=this_year).first()
                if bal:
                    bal.used_days += req.duration_days
                    bal.pending_days -= req.duration_days
                flash(f'Cuti {req.employee.name} disetujui ✅', 'success')
                company = Company.query.get(req.company_id)
                send_notification(company, 'approve', req)
            else:
                req.current_approval_level += 1
                req.status = 'pending'
                req.approved_by = None
                req.approved_at = None
                flash(f'Level {req.current_approval_level - 1} disetujui — menunggu approval level {req.current_approval_level}.', 'info')

        elif action == 'reject':
            req.status = 'rejected'
            req.notes = notes
            this_year = req.start_date.year
            bal = LeaveBalance.query.filter_by(employee_id=req.employee_id, leave_type_id=req.leave_type_id, year=this_year).first()
            if bal:
                bal.pending_days -= req.duration_days
            flash(f'Cuti {req.employee.name} ditolak ❌', 'warning')
            company = Company.query.get(req.company_id)
            send_notification(company, 'reject', req)

        else:
            flash('Aksi tidak valid.', 'danger')
            return redirect(url_for('main.approvals'))

        db.session.commit()
        return redirect(url_for('main.approvals'))

    # ── HISTORY ──

    @main_bp.route('/history')
    @login_required
    def history():
        year_filter = request.args.get('year', type=int)
        status_filter = request.args.get('status', '')
        this_year = date.today().year
        if not year_filter:
            year_filter = this_year

        years = list(range(this_year - 2, this_year + 3))
        my_company = current_user.company_id

        requests = LeaveRequest.query.filter(
            LeaveRequest.employee_id == current_user.id,
            LeaveRequest.company_id == my_company,
            db.extract('year', LeaveRequest.start_date) == year_filter
        )
        if status_filter:
            requests = requests.filter_by(status=status_filter)
        requests = requests.order_by(LeaveRequest.created_at.desc()).all()

        balances = LeaveBalance.query.filter_by(employee_id=current_user.id, year=year_filter).all()

        return render_template('history.html', requests=requests, years=years, year_filter=year_filter, status_filter=status_filter, balances=balances)

    # ── CALENDAR ──

    @main_bp.route('/calendar')
    @login_required
    def calendar():
        today = date.today()
        month = request.args.get('month', type=int, default=today.month)
        year = request.args.get('year', type=int, default=today.year)

        if month == 1:
            prev_month, prev_year = 12, year - 1
        else:
            prev_month, prev_year = month - 1, year
        if month == 12:
            next_month, next_year = 1, year + 1
        else:
            next_month, next_year = month + 1, year

        first_day = date(year, month, 1)
        if month == 12:
            last_day = date(year + 1, 1, 1) - timedelta(days=1)
        else:
            last_day = date(year, month + 1, 1) - timedelta(days=1)
        days_in_month = last_day.day
        first_weekday = first_day.weekday()

        my_company = current_user.company_id
        all_requests = LeaveRequest.query.filter(
            LeaveRequest.company_id == my_company,
            LeaveRequest.status.in_(['pending', 'approved']),
            LeaveRequest.start_date <= last_day,
            LeaveRequest.end_date >= first_day
        ).all()

        leaves_by_date = {}
        for req in all_requests:
            for d in range(max(req.start_date.day, 1), min(req.end_date.day, days_in_month) + 1):
                if d not in leaves_by_date:
                    leaves_by_date[d] = []
                leaves_by_date[d].append(req)

        month_names = ['', 'Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni',
                       'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']

        return render_template('calendar.html',
            month=month, year=year,
            prev_month=prev_month, prev_year=prev_year,
            next_month=next_month, next_year=next_year,
            month_name=month_names[month],
            days_in_month=days_in_month,
            first_weekday=first_weekday,
            leaves_by_date=leaves_by_date,
            today=today)

    # ── LONG WEEKEND ──

    @main_bp.route('/long-weekend')
    @login_required
    def long_weekend():
        year = request.args.get('year', type=int, default=date.today().year)
        years = list(range(2025, 2031))

        holidays = {
            2025: [(1,1), (1,27), (1,28), (1,29), (3,29), (3,30), (3,31), (4,1), (4,18), (5,1), (5,29), (6,6), (6,7), (6,8), (6,27), (8,17), (9,5), (12,25)],
            2026: [(1,1), (2,17), (2,18), (2,19), (3,20), (4,24), (5,1), (5,14), (5,22), (5,23), (5,24), (5,25), (7,8), (8,17), (12,25)],
            2027: [(1,1), (2,8), (2,9), (2,10), (3,11), (4,15), (5,1), (5,20), (5,21), (5,22), (5,23), (6,26), (8,17), (12,25)],
            2028: [(1,1), (1,27), (1,28), (1,29), (2,17), (3,31), (4,22), (5,1), (5,10), (5,11), (5,12), (5,13), (7,6), (8,17), (12,25)],
            2029: [(1,1), (2,15), (2,16), (2,17), (3,21), (4,5), (5,1), (5,30), (5,31), (6,1), (6,2), (6,25), (8,17), (12,25)],
            2030: [(1,1), (2,5), (2,6), (2,7), (3,21), (4,27), (5,1), (5,19), (5,20), (5,21), (5,22), (6,14), (8,17), (12,25)],
        }

        year_holidays = holidays.get(year, [])
        holiday_dates = [d(year, m, d) for m, d in year_holidays]

        suggestions = []
        for i in range(len(holiday_dates) - 1):
            h1 = holiday_dates[i]
            h2 = holiday_dates[i + 1]
            gap = (h2 - h1).days - 1
            if gap > 0:
                leave_needed = gap
                total_off = gap + 2
                suggestions.append({
                    'month': h1.strftime('%B'),
                    'leave_needed': leave_needed,
                    'total_off': total_off,
                    'detail': f'{h1.strftime("%d %b")} - {h2.strftime("%d %b")}'
                })

        return render_template('long_weekend.html', year=year, years=years, suggestions=suggestions)

    # ── ADMIN: COMPANIES ──

    @main_bp.route('/admin/companies')
    @login_required
    @role_required('admin')
    def admin_companies():
        query = Company.query.order_by(Company.name)
        
        # Search filter
        q = request.args.get('q', '').strip()
        if q:
            query = query.filter(Company.name.ilike(f'%{q}%'))
        
        # Notifications filter
        notif = request.args.get('notifications', '')
        if notif == 'enabled':
            query = query.filter_by(notifications_enabled=True)
        elif notif == 'disabled':
            query = query.filter_by(notifications_enabled=False)
        
        companies = query.all()
        return render_template('admin/companies.html', companies=companies)

    @main_bp.route('/admin/companies/add', methods=['POST'])
    @login_required
    @role_required('admin')
    def admin_add_company():
        name = request.form.get('name', '').strip()
        primary_color = request.form.get('primary_color', '#0d9488')
        if not name:
            flash('Nama perusahaan harus diisi.', 'danger')
            return redirect(url_for('main.admin_companies'))
        if Company.query.filter_by(name=name).first():
            flash('Perusahaan sudah ada.', 'danger')
            return redirect(url_for('main.admin_companies'))
        company = Company(name=name, primary_color=primary_color)
        db.session.add(company)
        db.session.commit()
        flash(f'Perusahaan {name} berhasil ditambahkan!', 'success')
        return redirect(url_for('main.admin_companies'))

    @main_bp.route('/admin/companies/edit/<int:company_id>', methods=['POST'])
    @login_required
    @role_required('admin')
    def admin_edit_company(company_id):
        company = Company.query.get_or_404(company_id)
        company.name = request.form.get('name', company.name)
        company.primary_color = request.form.get('primary_color', company.primary_color)
        company.notifications_enabled = bool(request.form.get('notifications_enabled'))
        company.smtp_host = request.form.get('smtp_host') or None
        company.smtp_port = request.form.get('smtp_port', type=int) or 587
        company.smtp_user = request.form.get('smtp_user') or None
        if request.form.get('smtp_password'):
            company.smtp_password = request.form.get('smtp_password')
        company.notify_on_approve = bool(request.form.get('notify_on_approve'))
        company.notify_on_reject = bool(request.form.get('notify_on_reject'))
        company.notify_on_submit = bool(request.form.get('notify_on_submit'))
        db.session.commit()
        flash(f'Perusahaan {company.name} berhasil diperbarui!', 'success')
        return redirect(url_for('main.admin_companies'))

    # ── ADMIN: EMPLOYEES ──

    @main_bp.route('/admin/employees')
    @login_required
    @role_required('hr', 'admin')
    def admin_employees():
        my_company = current_user.company_id
        
        # Build query with filters
        query = User.query.filter_by(company_id=my_company)
        
        # Search filter
        q = request.args.get('q', '').strip()
        if q:
            query = query.filter(
                (User.name.ilike(f'%{q}%')) | 
                (User.email.ilike(f'%{q}%'))
            )
        
        # Role filter
        role = request.args.get('role', '')
        if role:
            query = query.filter_by(role=role)
        
        # Department filter
        department_id = request.args.get('department_id', type=int)
        if department_id:
            query = query.filter_by(department_id=department_id)
        
        # Status filter
        status = request.args.get('status', '')
        if status == 'active':
            query = query.filter_by(is_active=True)
        elif status == 'inactive':
            query = query.filter_by(is_active=False)
        
        employees = query.order_by(User.role.desc(), User.name).all()
        managers = User.query.filter(User.company_id == my_company, User.role.in_(['manager', 'hr', 'admin'])).all()
        departments = Department.query.filter_by(company_id=my_company).all()
        return render_template('admin/employees.html', employees=employees, managers=managers, departments=departments)

    @main_bp.route('/admin/employees/add', methods=['POST'])
    @login_required
    @role_required('hr', 'admin')
    def admin_add_employee():
        my_company = current_user.company_id
        name = request.form.get('name', '').strip()
        email = request.form.get('email', '').strip()
        password = request.form.get('password', '')
        role = request.form.get('role', 'employee')
        manager_id = request.form.get('manager_id', type=int)
        department_id = request.form.get('department_id', type=int)

        if not all([name, email, password]):
            flash('Harap isi nama, email, dan password.', 'danger')
            return redirect(url_for('main.admin_employees'))

        if User.query.filter_by(email=email).first():
            flash('Email sudah terdaftar.', 'danger')
            return redirect(url_for('main.admin_employees'))

        user = User(name=name, email=email, role=role, company_id=my_company)
        user.set_password(password)
        if role == 'employee' and manager_id:
            user.manager_id = manager_id
        if role == 'employee' and department_id:
            user.department_id = department_id
        db.session.add(user)
        db.session.flush()

        this_year = date.today().year
        for lt in LeaveType.query.filter(LeaveType.days_per_year > 0, LeaveType.company_id == my_company).all():
            bal = LeaveBalance(employee_id=user.id, leave_type_id=lt.id, year=this_year, total_days=lt.days_per_year)
            db.session.add(bal)

        db.session.commit()
        flash(f'Karyawan {name} berhasil ditambahkan!', 'success')
        return redirect(url_for('main.admin_employees'))

    @main_bp.route('/admin/employees/delete/<int:user_id>', methods=['POST'])
    @login_required
    @role_required('hr', 'admin')
    def admin_delete_employee(user_id):
        my_company = current_user.company_id
        user = User.query.get_or_404(user_id)
        if user.role == 'admin':
            flash('Tidak bisa menghapus admin.', 'danger')
            return redirect(url_for('main.admin_employees'))
        name = user.name
        LeaveRequest.query.filter_by(employee_id=user.id).delete()
        LeaveBalance.query.filter_by(employee_id=user.id).delete()
        db.session.delete(user)
        db.session.commit()
        flash(f'Karyawan {name} berhasil dihapus.', 'success')
        return redirect(url_for('main.admin_employees'))

    # ── ADMIN: LEAVE TYPES ──

    @main_bp.route('/admin/leave-types')
    @login_required
    @role_required('hr', 'admin')
    def admin_leave_types():
        my_company = current_user.company_id
        
        # Build query with filters
        query = LeaveType.query.filter_by(company_id=my_company)
        
        # Search filter
        q = request.args.get('q', '').strip()
        if q:
            query = query.filter(LeaveType.name.ilike(f'%{q}%'))
        
        # Status filter
        status = request.args.get('status', '')
        if status == 'active':
            query = query.filter_by(is_active=True)
        elif status == 'inactive':
            query = query.filter_by(is_active=False)
        
        leave_types = query.order_by(LeaveType.name).all()
        return render_template('admin/leave_types.html', leave_types=leave_types)

    @main_bp.route('/admin/leave-types/add', methods=['POST'])
    @login_required
    @role_required('hr', 'admin')
    def admin_add_leave_type():
        my_company = current_user.company_id
        name = request.form.get('name', '').strip()
        days_per_year = request.form.get('days_per_year', type=int, default=0)
        color = request.form.get('color', '#0d9488')
        description = request.form.get('description', '')

        if not name:
            flash('Nama cuti harus diisi.', 'danger')
            return redirect(url_for('main.admin_leave_types'))

        lt = LeaveType(name=name, description=description, days_per_year=days_per_year, color=color, is_active=True, company_id=my_company)
        db.session.add(lt)
        db.session.flush()

        this_year = date.today().year
        for emp in User.query.filter(User.role == 'employee', User.company_id == my_company).all():
            bal = LeaveBalance(employee_id=emp.id, leave_type_id=lt.id, year=this_year, total_days=days_per_year)
            db.session.add(bal)
        db.session.commit()
        flash(f'Jenis cuti {name} berhasil ditambahkan!', 'success')
        return redirect(url_for('main.admin_leave_types'))

    @main_bp.route('/admin/leave-types/edit/<int:type_id>', methods=['POST'])
    @login_required
    @role_required('hr', 'admin')
    def admin_edit_leave_type(type_id):
        my_company = current_user.company_id
        lt = LeaveType.query.filter_by(id=type_id, company_id=my_company).first_or_404()
        lt.name = request.form.get('name', lt.name)
        lt.days_per_year = request.form.get('days_per_year', type=int, default=lt.days_per_year)
        lt.color = request.form.get('color', lt.color)
        lt.description = request.form.get('description', lt.description)
        lt.is_active = bool(request.form.get('is_active', lt.is_active))
        db.session.commit()
        flash(f'Jenis cuti {lt.name} berhasil diperbarui!', 'success')
        return redirect(url_for('main.admin_leave_types'))

    # ── ADMIN: DEPARTMENTS ──

    @main_bp.route('/admin/departments')
    @login_required
    @role_required('hr', 'admin')
    def admin_departments():
        my_company = current_user.company_id
        
        # Build query with filters
        query = Department.query.filter_by(company_id=my_company)
        
        # Search filter
        q = request.args.get('q', '').strip()
        if q:
            query = query.filter(Department.name.ilike(f'%{q}%'))
        
        # Head filter
        head_id = request.args.get('head_id', '')
        if head_id == 'none':
            query = query.filter(Department.head_id.is_(None))
        elif head_id:
            query = query.filter_by(head_id=int(head_id))
        
        departments = query.all()
        managers = User.query.filter(User.company_id == my_company, User.role.in_(['manager', 'hr', 'admin'])).all()
        return render_template('admin/departments.html', departments=departments, managers=managers)

    @main_bp.route('/admin/departments/add', methods=['POST'])
    @login_required
    @role_required('hr', 'admin')
    def admin_add_department():
        my_company = current_user.company_id
        name = request.form.get('name', '').strip()
        head_id = request.form.get('head_id', type=int)

        if not name:
            flash('Nama departemen harus diisi.', 'danger')
            return redirect(url_for('main.admin_departments'))

        dept = Department(name=name, company_id=my_company, head_id=head_id)
        db.session.add(dept)
        db.session.commit()
        flash(f'Departemen {name} berhasil ditambahkan!', 'success')
        return redirect(url_for('main.admin_departments'))

    @main_bp.route('/admin/departments/edit/<int:dept_id>', methods=['POST'])
    @login_required
    @role_required('hr', 'admin')
    def admin_edit_department(dept_id):
        dept = Department.query.get_or_404(dept_id)
        dept.name = request.form.get('name', dept.name)
        dept.head_id = request.form.get('head_id', type=int)
        db.session.commit()
        flash(f'Departemen {dept.name} berhasil diperbarui!', 'success')
        return redirect(url_for('main.admin_departments'))

    # ── ADMIN: APPROVAL CONFIGS ──

    @main_bp.route('/admin/approval-configs/<int:type_id>')
    @login_required
    @role_required('hr', 'admin')
    def admin_approval_configs(type_id):
        my_company = current_user.company_id
        leave_type = LeaveType.query.filter_by(id=type_id, company_id=my_company).first_or_404()
        configs = ApprovalConfig.query.filter(
            ApprovalConfig.company_id == my_company,
            (ApprovalConfig.leave_type_id == type_id) | (ApprovalConfig.leave_type_id.is_(None))
        ).order_by(ApprovalConfig.level).all()
        roles = [('manager', 'Manager'), ('hr', 'HR'), ('director', 'Direktur')]
        return render_template('admin/approval_configs.html', leave_type=leave_type, configs=configs, roles=roles)

    @main_bp.route('/admin/approval-configs/add/<int:type_id>', methods=['POST'])
    @login_required
    @role_required('hr', 'admin')
    def admin_add_approval_config(type_id):
        my_company = current_user.company_id
        level = request.form.get('level', type=int)
        approver_role = request.form.get('approver_role')

        if not level or not approver_role:
            flash('Level dan role harus diisi.', 'danger')
            return redirect(url_for('main.admin_approval_configs', type_id=type_id))

        config = ApprovalConfig(company_id=my_company, leave_type_id=type_id, level=level, approver_role=approver_role)
        db.session.add(config)
        db.session.commit()
        flash('Level approval berhasil ditambahkan!', 'success')
        return redirect(url_for('main.admin_approval_configs', type_id=type_id))

    @main_bp.route('/admin/approval-configs/delete/<int:config_id>', methods=['POST'])
    @login_required
    @role_required('hr', 'admin')
    def admin_delete_approval_config(config_id):
        config = ApprovalConfig.query.get_or_404(config_id)
        type_id = config.leave_type_id
        db.session.delete(config)
        db.session.commit()
        flash('Level approval berhasil dihapus!', 'success')
        return redirect(url_for('main.admin_approval_configs', type_id=type_id))

    # ── EXPORT: Excel Report ──

    @main_bp.route('/export/excel')
    @login_required
    @role_required('hr', 'admin')
    def export_excel():
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        my_company = current_user.company_id
        this_year = date.today().year

        requests = LeaveRequest.query.filter(
            LeaveRequest.company_id == my_company,
            db.extract('year', LeaveRequest.start_date) == this_year
        ).order_by(LeaveRequest.start_date).all()

        wb = Workbook()
        ws = wb.active
        ws.title = f'Cuti {this_year}'

        hdr_font = Font(bold=True, color='FFFFFF', size=11)
        hdr_fill = PatternFill(start_color='0d9488', end_color='0d9488', fill_type='solid')
        headers = ['No', 'Karyawan', 'Departemen', 'Jenis Cuti', 'Tanggal Mulai', 'Tanggal Selesai', 'Durasi (hari)', 'Status', 'Alasan', 'Delegate', 'Disetujui Oleh']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal='center')

        for i, req in enumerate(requests, 1):
            row = i + 1
            ws.cell(row=row, column=1, value=i)
            ws.cell(row=row, column=2, value=req.employee.name)
            ws.cell(row=row, column=3, value=req.employee.department.name if req.employee.department else '-')
            ws.cell(row=row, column=4, value=req.leave_type.name)
            ws.cell(row=row, column=5, value=req.start_date.strftime('%d/%m/%Y'))
            ws.cell(row=row, column=6, value=req.end_date.strftime('%d/%m/%Y'))
            ws.cell(row=row, column=7, value=int(req.duration_days))
            ws.cell(row=row, column=8, value=req.status.title())
            ws.cell(row=row, column=9, value=req.reason or '-')
            ws.cell(row=row, column=10, value=req.delegate.name if req.delegate else '-')
            ws.cell(row=row, column=11, value=req.approver.name if req.approver else '-')

        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(output, download_name=f'gwcuti-report-{this_year}.xlsx', as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    # ── API: Chart Data ──

    @main_bp.route('/api/chart-data')
    @login_required
    def api_chart_data():
        this_year = date.today().year
        months = ['Jan','Feb','Mar','Apr','Mei','Jun','Jul','Agu','Sep','Okt','Nov','Des']

        if current_user.role == 'admin':
            company_id = request.args.get('company_id', type=int)
            if company_id:
                my_company = company_id
            else:
                my_company = current_user.company_id
        else:
            my_company = current_user.company_id

        if current_user.role == 'employee':
            base = LeaveRequest.query.filter_by(employee_id=current_user.id, company_id=my_company)
        else:
            base = LeaveRequest.query.filter_by(company_id=my_company)

        monthly = []
        for m in range(1, 13):
            count = base.filter(
                db.extract('year', LeaveRequest.start_date) == this_year,
                db.extract('month', LeaveRequest.start_date) == m
            ).count()
            monthly.append(count)

        types = LeaveType.query.filter_by(company_id=my_company).all()
        type_labels = [t.name for t in types]
        type_counts = []
        for t in types:
            count = base.filter(LeaveRequest.leave_type_id == t.id).count()
            type_counts.append(count)

        companies = []
        if current_user.role == 'admin':
            companies = [{'id': c.id, 'name': c.name} for c in Company.query.filter_by(is_active=True).all()]

        return jsonify({
            'months': months,
            'monthly': monthly,
            'type_labels': type_labels,
            'type_counts': type_counts,
            'companies': companies,
            'current_company': my_company
        })

    app.register_blueprint(main_bp)


if __name__ == '__main__':
    app = create_app()
    app.run(host='0.0.0.0', port=5000, debug=True)
